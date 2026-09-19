import os
import shutil
import requests

from flask import Flask, render_template, request, jsonify
from dotenv import load_dotenv

from pypdf import PdfReader
from docx import Document as DocxDocument

from langchain_core.documents import Document
from langchain_text_splitters import RecursiveCharacterTextSplitter
from langchain_huggingface import HuggingFaceEmbeddings
from langchain_chroma import Chroma


# --------------------------------------------------
# LOAD ENVIRONMENT VARIABLES
# --------------------------------------------------

load_dotenv()

OLLAMA_URL = os.getenv(
    "OLLAMA_URL",
    "http://localhost:11434"
)

OLLAMA_MODEL = os.getenv(
    "OLLAMA_MODEL",
    "llama3.2"
)

EMBEDDING_MODEL = os.getenv(
    "EMBEDDING_MODEL",
    "all-MiniLM-L6-v2"
)


# --------------------------------------------------
# FLASK APP
# --------------------------------------------------

app = Flask(__name__)


# --------------------------------------------------
# FOLDERS
# --------------------------------------------------

DOCUMENT_FOLDER = "documents"
CHROMA_FOLDER = "chroma_db"

os.makedirs(DOCUMENT_FOLDER, exist_ok=True)
os.makedirs(CHROMA_FOLDER, exist_ok=True)


# --------------------------------------------------
# EMBEDDING MODEL
# --------------------------------------------------

print("Loading embedding model...")

embeddings = HuggingFaceEmbeddings(
    model_name=EMBEDDING_MODEL
)

print("Embedding model loaded.")


# --------------------------------------------------
# HOME PAGE
# --------------------------------------------------

@app.route("/")
def home():
    return render_template("index.html")


# --------------------------------------------------
# PDF LOADER
# --------------------------------------------------

def load_pdf(file_path):

    documents = []

    reader = PdfReader(file_path)

    for page_number, page in enumerate(reader.pages, start=1):

        text = page.extract_text()

        if text and text.strip():

            document = Document(
                page_content=text,
                metadata={
                    "source": os.path.basename(file_path),
                    "page": page_number
                }
            )

            documents.append(document)

    return documents


# --------------------------------------------------
# DOCX LOADER
# --------------------------------------------------

def load_docx(file_path):

    documents = []

    doc = DocxDocument(file_path)

    text_parts = []

    for paragraph in doc.paragraphs:

        text = paragraph.text.strip()

        if text:
            text_parts.append(text)

    full_text = "\n".join(text_parts)

    if full_text.strip():

        document = Document(
            page_content=full_text,
            metadata={
                "source": os.path.basename(file_path),
                "page": "N/A"
            }
        )

        documents.append(document)

    return documents


# --------------------------------------------------
# TXT LOADER
# --------------------------------------------------

def load_txt(file_path):

    documents = []

    with open(
        file_path,
        "r",
        encoding="utf-8",
        errors="ignore"
    ) as file:

        text = file.read()

    if text.strip():

        document = Document(
            page_content=text,
            metadata={
                "source": os.path.basename(file_path),
                "page": "N/A"
            }
        )

        documents.append(document)

    return documents


# --------------------------------------------------
# DOCUMENT LOADER
# --------------------------------------------------

def load_document(file_path):

    extension = os.path.splitext(file_path)[1].lower()

    if extension == ".pdf":

        return load_pdf(file_path)

    elif extension == ".docx":

        return load_docx(file_path)

    elif extension == ".txt":

        return load_txt(file_path)

    else:

        return []


# --------------------------------------------------
# TEXT CLEANING
# --------------------------------------------------

def clean_text(text):

    text = text.replace("\x00", " ")

    text = " ".join(text.split())

    return text.strip()


# --------------------------------------------------
# CLEAN ALL DOCUMENTS
# --------------------------------------------------

def clean_documents(documents):

    cleaned_documents = []

    for document in documents:

        cleaned_text = clean_text(
            document.page_content
        )

        if cleaned_text:

            document.page_content = cleaned_text

            cleaned_documents.append(document)

    return cleaned_documents


# --------------------------------------------------
# TEXT CHUNKING
# --------------------------------------------------

def split_documents(documents):

    text_splitter = RecursiveCharacterTextSplitter(
        chunk_size=800,
        chunk_overlap=150,
        separators=[
            "\n\n",
            "\n",
            ". ",
            " ",
            ""
        ]
    )

    chunks = text_splitter.split_documents(
        documents
    )

    return chunks


# --------------------------------------------------
# CREATE / LOAD CHROMA DATABASE
# --------------------------------------------------

def get_vector_database():

    vector_db = Chroma(
        collection_name="documind_documents",
        embedding_function=embeddings,
        persist_directory=CHROMA_FOLDER
    )

    return vector_db


# --------------------------------------------------
# ADD DOCUMENTS TO CHROMA
# --------------------------------------------------

def add_to_vector_database(chunks):

    vector_db = get_vector_database()

    vector_db.add_documents(chunks)

    return vector_db


# --------------------------------------------------
# CHECK OLLAMA
# --------------------------------------------------

def check_ollama():

    try:

        response = requests.get(
            f"{OLLAMA_URL}/api/tags",
            timeout=5
        )

        if response.status_code == 200:

            return True

        return False

    except Exception:

        return False


# --------------------------------------------------
# ASK LOCAL OLLAMA
# --------------------------------------------------

def ask_ollama(question, context):

    prompt = f"""
You are DocuMind AI, an intelligent document question-answering assistant.

You must answer the user's question ONLY using the information
provided in the DOCUMENT CONTEXT.

Do not use outside knowledge.

Do not guess.

Do not invent information.

If the answer is not available in the document context,
reply exactly:

The information was not found in the uploaded documents.

Keep the answer clear and easy to understand.

DOCUMENT CONTEXT:
-----------------
{context}
-----------------

USER QUESTION:
{question}

ANSWER:
"""

    try:

        response = requests.post(
            f"{OLLAMA_URL}/api/generate",
            json={
                "model": OLLAMA_MODEL,
                "prompt": prompt,
                "stream": False
            },
            timeout=180
        )

        if response.status_code != 200:

            return None, "Ollama returned an error."

        data = response.json()

        answer = data.get("response", "").strip()

        if not answer:

            return None, "No answer was generated."

        return answer, None

    except requests.exceptions.ConnectionError:

        return None, (
            "Ollama is not running. "
            "Please start Ollama and try again."
        )

    except requests.exceptions.Timeout:

        return None, (
            "Ollama took too long to respond. "
            "Please try again."
        )

    except Exception as error:

        return None, str(error)


# --------------------------------------------------
# UPLOAD DOCUMENT
# --------------------------------------------------

@app.route("/upload", methods=["POST"])
def upload_document():

    try:

        if "file" not in request.files:

            return jsonify({
                "success": False,
                "message": "No file selected."
            }), 400

        file = request.files["file"]

        if file.filename == "":

            return jsonify({
                "success": False,
                "message": "Please select a file."
            }), 400

        allowed_extensions = {
            ".pdf",
            ".docx",
            ".txt"
        }

        extension = os.path.splitext(
            file.filename
        )[1].lower()

        if extension not in allowed_extensions:

            return jsonify({
                "success": False,
                "message": (
                    "Only PDF, DOCX and TXT files "
                    "are supported."
                )
            }), 400

        filename = os.path.basename(
            file.filename
        )

        file_path = os.path.join(
            DOCUMENT_FOLDER,
            filename
        )

        file.save(file_path)

        # Load document
        documents = load_document(file_path)

        if not documents:

            os.remove(file_path)

            return jsonify({
                "success": False,
                "message": (
                    "No readable text was found "
                    "in the document."
                )
            }), 400

        # Clean text
        documents = clean_documents(
            documents
        )

        # Split into chunks
        chunks = split_documents(
            documents
        )

        if not chunks:

            return jsonify({
                "success": False,
                "message": (
                    "Could not create document chunks."
                )
            }), 400

        # Add chunks to ChromaDB
        add_to_vector_database(
            chunks
        )

        return jsonify({
            "success": True,
            "message": (
                f"{filename} uploaded successfully."
            ),
            "chunks": len(chunks)
        })

    except Exception as error:

        return jsonify({
            "success": False,
            "message": str(error)
        }), 500


# --------------------------------------------------
# ASK QUESTION
# --------------------------------------------------

@app.route("/ask", methods=["POST"])
def ask_question():

    try:

        data = request.get_json()

        question = data.get(
            "question",
            ""
        ).strip()

        if not question:

            return jsonify({
                "success": False,
                "message": "Please enter a question."
            }), 400

        # Check if documents exist
        files = os.listdir(
            DOCUMENT_FOLDER
        )

        if not files:

            return jsonify({
                "success": False,
                "message": (
                    "Please upload a document first."
                )
            }), 400

        # Check Ollama
        if not check_ollama():

            return jsonify({
                "success": False,
                "message": (
                    "Ollama is not running. "
                    "Please start Ollama first."
                )
            }), 500

        # Load ChromaDB
        vector_db = get_vector_database()

        # Similarity search
        relevant_documents = vector_db.similarity_search(
            question,
            k=4
        )

        if not relevant_documents:

            return jsonify({
                "success": True,
                "answer": (
                    "The information was not found "
                    "in the uploaded documents."
                ),
                "sources": []
            })

        # Create context
        context_parts = []

        sources = []

        for document in relevant_documents:

            source = document.metadata.get(
                "source",
                "Unknown"
            )

            page = document.metadata.get(
                "page",
                "N/A"
            )

            context_parts.append(
                document.page_content
            )

            sources.append({
                "source": source,
                "page": page
            })

        context = "\n\n".join(
            context_parts
        )

        # Ask Ollama
        answer, error = ask_ollama(
            question,
            context
        )

        if error:

            return jsonify({
                "success": False,
                "message": error
            }), 500

        # Remove duplicate sources
        unique_sources = []

        seen = set()

        for source in sources:

            key = (
                source["source"],
                source["page"]
            )

            if key not in seen:

                seen.add(key)

                unique_sources.append(
                    source
                )

        return jsonify({
            "success": True,
            "answer": answer,
            "sources": unique_sources
        })

    except Exception as error:

        return jsonify({
            "success": False,
            "message": str(error)
        }), 500


# --------------------------------------------------
# CLEAR DOCUMENTS
# --------------------------------------------------

@app.route("/clear", methods=["POST"])
def clear_documents():

    try:

        # Delete uploaded files
        if os.path.exists(
            DOCUMENT_FOLDER
        ):

            for filename in os.listdir(
                DOCUMENT_FOLDER
            ):

                file_path = os.path.join(
                    DOCUMENT_FOLDER,
                    filename
                )

                if os.path.isfile(file_path):

                    os.remove(file_path)

        # Delete ChromaDB
        if os.path.exists(
            CHROMA_FOLDER
        ):

            shutil.rmtree(
                CHROMA_FOLDER
            )

        # Recreate empty Chroma folder
        os.makedirs(
            CHROMA_FOLDER,
            exist_ok=True
        )

        return jsonify({
            "success": True,
            "message": "All documents cleared."
        })

    except Exception as error:

        return jsonify({
            "success": False,
            "message": str(error)
        }), 500


# --------------------------------------------------
# RUN APPLICATION
# --------------------------------------------------

if __name__ == "__main__":

    app.run(
        host="127.0.0.1",
        port=5000,
        debug=True
    )